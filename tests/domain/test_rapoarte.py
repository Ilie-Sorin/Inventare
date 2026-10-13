import csv
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from inventar.domain.comun import EroareValidare
from inventar.domain.generare import Perimetru, genereaza
from inventar.models import Jurnal
from inventar.rapoarte import centralizator_fisier, export_linii
from inventar.tiparire import tipareste_declaratii, tipareste_liste
from tests.conftest import CONSULTARE, OPERATOR


@pytest.fixture
def inv(session, persoana, bunuri, initial, nom):
    a = persoana("ANA ŢUŢUIANU")
    b = persoana("BOGDAN", locatie="trezorerie_loc", activitate="trezorerie_act")
    [m1] = bunuri("1", cont="213.01", pret="1234.50", data_amortizare_integrala=date(2030, 1, 31))
    [m2] = bunuri("2", cont="213.01", pret="100.00")
    [o1] = bunuri("3", categorie="OI", cont="303.02", pret="10.25")
    initial(date(2025, 1, 1), [(m1, a), (m2, b), (o1, a)])
    return {"inv": genereaza(session, OPERATOR, date(2025, 11, 28), Perimetru()), "a": a, "b": b, "nom": nom}


def test_export_xlsx(session, inv):
    nume, continut, mime = export_linii(session, CONSULTARE, inv["inv"].id, "xlsx")
    assert nume == "Export_inventar_2025-11-28.xlsx" and "spreadsheet" in mime
    foaie = load_workbook(BytesIO(continut)).active
    randuri = list(foaie.iter_rows(values_only=True))
    assert randuri[0][:3] == ("Nr. inventar", "Nr. inventar vechi", "Denumire")
    assert len(randuri) == 4
    primul = next(r for r in randuri[1:] if r[0] == "1")
    assert primul[8] == pytest.approx(1234.5) and primul[10] == "ANA ŢUŢUIANU"
    assert primul[9].date() == date(2030, 1, 31)


def test_export_csv_filtrat(session, inv):
    filtru = Perimetru(locatii=(inv["nom"]["trezorerie_loc"].id,))
    _, continut, _ = export_linii(session, OPERATOR, inv["inv"].id, "csv", filtru)
    assert continut.startswith("﻿".encode("utf-8"))
    randuri = list(csv.reader(StringIO(continut.decode("utf-8-sig")), delimiter=";"))
    assert len(randuri) == 2 and randuri[1][0] == "2" and randuri[1][8] == "100,00"


def test_centralizator_pdf_si_xlsx(session, inv):
    nume, continut, _ = centralizator_fisier(session, OPERATOR, inv["inv"].id, "categorii", "pdf")
    assert nume == "Centralizator_categorii_2025-11-28.pdf"
    text = PdfReader(BytesIO(continut)).pages[0].extract_text()
    for asteptat in ("Centralizator pe categorii", "TOTAL GENERAL", "1.344,75", "Pagina 1 din 1", "Filtru: Tot"):
        assert asteptat in text, asteptat
    _, continut, _ = centralizator_fisier(session, OPERATOR, inv["inv"].id, "gestiuni", "xlsx",
                                          Perimetru(categorii=("MF",)))
    valori = [r for r in load_workbook(BytesIO(continut)).active.iter_rows(values_only=True)]
    assert valori[1][0] == "Filtru: Categorii: MF"
    assert valori[-1][:2] == ("TOTAL GENERAL", 2) and Decimal(str(valori[-1][2])) == Decimal("1334.5")


def test_listele_filtrate(session, inv):
    nume, continut = tipareste_liste(session, OPERATOR, inv["inv"].id, "MF", None,
                                     filtru=Perimetru(locatii=(inv["nom"]["botosani"].id,)))
    assert nume == "Inventar_2025-11-28_selectie_MF.pdf"
    text = PdfReader(BytesIO(continut)).pages[0].extract_text()
    assert "ANA ŢUŢUIANU" in text and "BOGDAN" not in text
    # Un filtru pe categorie care exclude categoria cerută nu produce nimic.
    with pytest.raises(EroareValidare, match="Nu există bunuri"):
        tipareste_declaratii(session, OPERATOR, inv["inv"].id, "MF", "inceput", filtru=Perimetru(categorii=("OI",)))


def test_neactualizat_nu_se_exporta(session, inv):
    inv["inv"].stare = "NEACTUALIZAT"
    session.commit()
    with pytest.raises(EroareValidare, match="regenerați"):
        export_linii(session, OPERATOR, inv["inv"].id, "csv")


def test_jurnal(session, inv):
    export_linii(session, OPERATOR, inv["inv"].id, "csv")
    centralizator_fisier(session, OPERATOR, inv["inv"].id, "conturi", "xlsx")
    actiuni = list(session.scalars(select(Jurnal.actiune).where(Jurnal.actiune.in_(["EXPORT", "CENTRALIZATOR"]))
                                   .order_by(Jurnal.id)))
    assert actiuni == ["EXPORT", "CENTRALIZATOR"]
