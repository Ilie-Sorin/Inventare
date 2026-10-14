from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from inventar.domain.comun import EroareValidare
from inventar.drepturi import AccesInterzis
from inventar.importuri.email_persoane import analizeaza, aplica, exporta
from inventar.models import Jurnal
from tests.conftest import CONSULTARE, OPERATOR


def _xlsx(randuri: list[list], antet=("ID", "Gestiune", "Titular", "E-mail", "Activ")) -> bytes:
    carte = Workbook()
    foaie = carte.active
    foaie.append(list(antet))
    for rand in randuri:
        foaie.append(rand)
    iesire = BytesIO()
    carte.save(iesire)
    return iesire.getvalue()


@pytest.fixture
def lume(persoana):
    ana = persoana("ANA", email="ana@vechi.ro")
    dan = persoana("DAN")
    dan_et2 = persoana("DAN ET2", titular_id=dan.id)
    return {"ana": ana, "dan": dan, "dan_et2": dan_et2}


def test_export_contine_gestiunile_active_cu_id_si_email(session, lume):
    inactiv = lume["dan_et2"]
    inactiv.activ = False
    session.commit()

    foaie = load_workbook(BytesIO(exporta(session, CONSULTARE))).active
    randuri = list(foaie.iter_rows(values_only=True))
    assert randuri[0] == ("ID", "Gestiune", "Titular", "E-mail", "Activ")
    assert (lume["ana"].id, "ANA", None, "ana@vechi.ro", "da") in [tuple(r) for r in randuri[1:]]
    assert {r[1] for r in randuri[1:]} == {"ANA", "DAN"}

    complet = load_workbook(BytesIO(exporta(session, CONSULTARE, include_inactive=True))).active
    assert {r[1] for r in list(complet.iter_rows(values_only=True))[1:]} == {"ANA", "DAN", "DAN ET2"}
    assert session.scalars(select(Jurnal).where(Jurnal.actiune == "EXPORT_EMAIL_PERSOANE")).all()


def test_analiza_clasifica_randurile(session, lume):
    ana, dan = lume["ana"], lume["dan"]
    continut = _xlsx([
        [ana.id, "ANA", "", "ana@nou.ro", "da"],        # schimbat
        [dan.id, "DAN", "", "Dan@Exemplu.ro", "da"],     # nou
        [lume["dan_et2"].id, "DAN ET2", "", "", "da"],   # gol → neschimbat
        [ana.id, "ANA", "", "alta@x.ro", "da"],          # ID dublat
        [9999, "NECUNOSCUT", "", "x@y.ro", "da"],        # ID inexistent
        [None, "FĂRĂ ID", "", "x@y.ro", "da"],           # ID lipsă
        [str(dan.id) + ".0", "DAN", "", "nu-este-adresa", "da"],  # ID ca text, dublat (DAN apare mai sus)
    ])
    a = analizeaza(session, OPERATOR, continut)
    assert [(s.persoana_id, s.vechi, s.nou) for s in a.schimbari] == [
        (ana.id, "ana@vechi.ro", "ana@nou.ro"), (dan.id, None, "Dan@Exemplu.ro")]
    assert a.goale == 1 and a.neschimbate == 0
    assert len(a.erori) == 4
    assert any("de două ori" in e for e in a.erori)
    assert any("9999" in e for e in a.erori)
    assert any("ID lipsă" in e for e in a.erori)
    # Analiza nu modifică nimic.
    assert ana.email == "ana@vechi.ro" and dan.email is None


def test_adresa_identica_chiar_cu_alte_majuscule_e_neschimbata(session, lume):
    a = analizeaza(session, OPERATOR, _xlsx([[lume["ana"].id, "ANA", "", " ANA@vechi.RO ", "da"]]))
    assert (len(a.schimbari), a.neschimbate, a.erori) == (0, 1, [])


def test_adresa_invalida_e_eroare(session, lume):
    a = analizeaza(session, OPERATOR, _xlsx([[lume["ana"].id, "ANA", "", "fara-arond", "da"]]))
    assert a.schimbari == [] and "nu este validă" in a.erori[0]


def test_avertisment_pentru_aceeasi_adresa_la_persoane_diferite(session, lume):
    a = analizeaza(session, OPERATOR, _xlsx([[lume["dan"].id, "DAN", "", "ana@vechi.ro", "da"]]))
    assert len(a.schimbari) == 1
    assert "ANA" in a.avertismente[0] and "DAN" in a.avertismente[0]


def test_acelasi_titular_nu_da_avertisment(session, lume):
    a = analizeaza(session, OPERATOR, _xlsx([
        [lume["dan"].id, "DAN", "", "dan@x.ro", "da"], [lume["dan_et2"].id, "DAN ET2", "DAN", "dan@x.ro", "da"]]))
    assert len(a.schimbari) == 2 and a.avertismente == []


def test_fisier_fara_coloanele_cerute_sau_nevalid(session, lume):
    with pytest.raises(EroareValidare, match="coloanele"):
        analizeaza(session, OPERATOR, _xlsx([[1, "x"]], antet=("Nr", "Nume")))
    with pytest.raises(EroareValidare, match="nu este un document Excel"):
        analizeaza(session, OPERATOR, b"nu e excel")


def test_aplicare_modifica_doar_emailul_si_scrie_in_jurnal(session, lume):
    ana, dan = lume["ana"], lume["dan"]
    nume_ana = ana.nume
    n = aplica(session, OPERATOR, {ana.id: "ana@nou.ro", dan.id: "dan@x.ro"})
    assert n == 2
    assert (ana.email, ana.nume, dan.email) == ("ana@nou.ro", nume_ana, "dan@x.ro")
    intrare = session.scalars(select(Jurnal).where(Jurnal.actiune == "EMAIL_PERSOANE_IMPORTAT")).one()
    assert '"nr": 2' in intrare.detalii_json and "ana@vechi.ro" in intrare.detalii_json


def test_aplicarea_e_atomica(session, lume):
    ana = lume["ana"]
    with pytest.raises(EroareValidare):
        aplica(session, OPERATOR, {ana.id: "ana@nou.ro", 9999: "x@y.ro"})
    session.rollback()
    assert ana.email == "ana@vechi.ro"
    with pytest.raises(EroareValidare, match="nu este validă"):
        aplica(session, OPERATOR, {ana.id: "gresit"})


def test_consultare_nu_poate_importa(session, lume):
    with pytest.raises(AccesInterzis):
        analizeaza(session, CONSULTARE, _xlsx([]))
    with pytest.raises(AccesInterzis):
        aplica(session, CONSULTARE, {lume["ana"].id: "a@b.ro"})
