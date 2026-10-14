import json
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from inventar.models import Persoana


def _xlsx(randuri: list[list]) -> BytesIO:
    carte = Workbook()
    carte.active.append(["ID", "Gestiune", "Titular", "E-mail", "Activ"])
    for rand in randuri:
        carte.active.append(rand)
    iesire = BytesIO()
    carte.save(iesire)
    iesire.seek(0)
    return iesire


@pytest.fixture
def lume(client, creeaza, persoana):
    ana = persoana("ANA", email="ana@vechi.ro")
    dan = persoana("DAN")
    creeaza("op", rol="OPERATOR")
    client.post("/autentificare", data={"username": "op", "parola": "parola123"})
    return {"ana": ana, "dan": dan}


def _incarca(client, randuri):
    return client.post("/persoane/email/previzualizare", content_type="multipart/form-data",
                       data={"fisier": (_xlsx(randuri), "emailuri.xlsx")})


def test_export_si_pagina(client, lume):
    assert "E-mailuri" in client.get("/persoane/email").get_data(as_text=True)
    r = client.get("/persoane/email/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.mimetype
    foaie = load_workbook(BytesIO(r.data)).active
    assert [c.value for c in foaie[2]][:2] == [lume["ana"].id, "ANA"]


def test_flux_import_cu_previzualizare_si_confirmare(client, session, lume):
    ana, dan = lume["ana"], lume["dan"]
    r = _incarca(client, [[ana.id, "ANA", "", "ana@nou.ro", "da"], [dan.id, "DAN", "", "gresit", "da"]])
    pagina = r.get_data(as_text=True)
    assert r.status_code == 200
    assert "ana@nou.ro" in pagina and "nu este validă" in pagina
    assert ana.email == "ana@vechi.ro"  # doar previzualizare, nimic salvat

    schimbari = json.dumps({ana.id: "ana@nou.ro"})
    r = client.post("/persoane/email/aplicare", data={"schimbari": schimbari})
    assert r.status_code == 302
    assert session.get(Persoana, ana.id).email == "ana@nou.ro"
    assert session.get(Persoana, dan.id).email is None


def test_fisier_gresit_nu_cade(client, lume):
    r = client.post("/persoane/email/previzualizare", content_type="multipart/form-data",
                    data={"fisier": (BytesIO(b"nu e excel"), "x.xlsx")})
    assert r.status_code == 200 and "nu este un document Excel" in r.get_data(as_text=True)


def test_date_de_confirmare_stricate(client, lume):
    r = client.post("/persoane/email/aplicare", data={"schimbari": "nu-e-json"})
    assert r.status_code == 302


def test_consultare_exporta_dar_nu_importa(client, creeaza, lume):
    client.post("/deconectare")
    creeaza("cons", rol="CONSULTARE")
    client.post("/autentificare", data={"username": "cons", "parola": "parola123"})
    assert client.get("/persoane/email/export.xlsx").status_code == 200
    assert _incarca(client, []).status_code == 403
    assert client.post("/persoane/email/aplicare", data={"schimbari": "{}"}).status_code == 403
