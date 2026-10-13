from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from inventar.importuri.mf_baza import RezultatCitire
from inventar.models import Bun, Sincronizare
from tests.importuri.test_import_initial import bun as rand_export


@pytest.fixture
def lume(client, session, creeaza, persoana, bunuri, initial, monkeypatch):
    a = persoana("ANA")
    [b1] = bunuri("1", cont="213.01", pret="100.00")
    initial(date(2025, 1, 1), [(b1, a)])
    export = [rand_export("1", pret="150.00", cont="213.01"), rand_export("9", cont="214.01", pret="20.00")]
    monkeypatch.setattr("inventar.web.sincronizari.citeste_mf_baza",
                        lambda continut, fisier: RezultatCitire(list(export), []))
    creeaza("op", rol="OPERATOR")
    client.post("/autentificare", data={"username": "op", "parola": "parola123"})
    return {"b1": b1}


def _incarca(client):
    return client.post("/sincronizari/incarcare", data={"mf_baza": [(BytesIO(b"xls"), "MIJLOACE FIXE.xls")]},
                       content_type="multipart/form-data")


def test_flux_sincronizare(client, session, lume):
    r = _incarca(client)
    assert r.status_code == 302
    url = r.headers["Location"]
    pagina = client.get(url).get_data(as_text=True)
    assert "Bunuri noi în export" in pagina and "valoare: 100.00 → 150.00" in pagina
    assert "MIJLOACE FIXE.xls" in pagina
    assert lume["b1"].pret == Decimal("100.00")  # încă nimic aplicat

    s = session.scalar(select(Sincronizare))
    nou = next(d for d in s.diferente if d.tip == "BUN_NOU")
    client.post(f"{url}/aplicare", data={"diferente": [nou.id]})
    assert s.stare == "APLICATA"
    assert session.scalar(select(Bun).where(Bun.nr_inventar == "9")) is not None
    assert lume["b1"].pret == Decimal("100.00")  # nebifat → neaplicat
    pagina = client.get(url).get_data(as_text=True)
    assert "nerepartizate" in pagina and "neales" in pagina

    r = client.get(f"{url}/raport.xlsx")
    foi = load_workbook(BytesIO(r.data))
    assert foi.sheetnames == ["Totaluri", "Diferențe"]
    assert {rand[0] for rand in foi["Diferențe"].iter_rows(min_row=2, values_only=True)} == {"BUN_NOU", "MODIFICAT"}
    r = client.get(f"{url}/raport.pdf")
    assert "Sincronizarea nr." in PdfReader(BytesIO(r.data)).pages[0].extract_text()


def test_abandonare(client, session, lume):
    url = _incarca(client).headers["Location"]
    client.post(f"{url}/abandonare")
    assert session.scalar(select(Sincronizare.stare)) == "ABANDONATA"
    r = client.post(f"{url}/aplicare", data={}, follow_redirects=True)
    assert "deja aplicată sau abandonată" in r.get_data(as_text=True)


def test_consultarea_doar_vede(client, session, lume, creeaza):
    url = _incarca(client).headers["Location"]
    client.post("/deconectare")
    creeaza("cons", rol="CONSULTARE")
    client.post("/autentificare", data={"username": "cons", "parola": "parola123"})
    assert client.get("/sincronizari").status_code == 200
    assert "Încarcă și compară" not in client.get("/sincronizari").get_data(as_text=True)
    assert client.get(url).status_code == 200
    assert _incarca(client).status_code == 403
    assert client.post(f"{url}/aplicare").status_code == 403
