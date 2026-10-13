from datetime import date, datetime

import pytest
from sqlalchemy import select

from inventar.models import Document, Inventar

START = date(2025, 1, 1)


@pytest.fixture
def lume(client, session, creeaza, persoana, bunuri, initial):
    a = persoana("HRITCU MARCELA")
    b1, b2 = bunuri("1", "2", pret="100.00")
    [b3] = bunuri("3", categorie="OI", cont="303.02", pret="10.00")
    initial(START, [(b1, a), (b2, a), (b3, a)])
    # Exportul „încărcat” înaintea datelor din teste, ca să apară avertismentul.
    session.scalar(select(Document)).creat_la = datetime(2025, 1, 1)
    session.commit()
    creeaza("op", rol="OPERATOR")
    creeaza("admin")
    client.post("/autentificare", data={"username": "op", "parola": "parola123"})
    return {"a": a}


def test_generare_cu_avertismente_si_confirmare(client, session, lume):
    r = client.post("/inventare/nou", data={"data": "2025-11-28"})
    pagina = r.get_data(as_text=True)
    assert "Verificați înainte de generare" in pagina and "Generează oricum" in pagina
    assert session.scalar(select(Inventar)) is None

    r = client.post("/inventare/nou", data={"data": "2025-11-28", "confirmat": "1"})
    assert r.status_code == 302
    inventar = session.scalar(select(Inventar))
    pagina = client.get(f"/inventare/{inventar.id}").get_data(as_text=True)
    assert "Liste (2)" in pagina and "HRITCU MARCELA" in pagina and "200,00" in pagina
    assert "Tot" in client.get("/inventare").get_data(as_text=True)


def test_generare_cu_perimetru(client, session, lume):
    client.post("/inventare/nou", data={"data": "2025-11-28", "categorii": ["OI"], "confirmat": "1"})
    inventar = session.scalar(select(Inventar))
    assert inventar.nr_linii == 1
    assert "Categorii: OI" in client.get("/inventare").get_data(as_text=True)


def test_blocare_si_deblocare_doar_admin(client, session, lume):
    client.post("/inventare/nou", data={"data": "2025-11-28", "confirmat": "1"})
    inventar = session.scalar(select(Inventar))
    client.post(f"/inventare/{inventar.id}/blocare")
    assert inventar.stare == "BLOCAT"
    r = client.post(f"/inventare/{inventar.id}/stergere", follow_redirects=True)
    assert "blocat" in r.get_data(as_text=True)
    assert client.post(f"/inventare/{inventar.id}/deblocare", data={"motiv": "x"}).status_code == 403

    client.post("/deconectare")
    client.post("/autentificare", data={"username": "admin", "parola": "parola123"})
    client.post(f"/inventare/{inventar.id}/deblocare", data={"motiv": "corectură"})
    assert inventar.stare == "NEACTUALIZAT"
    client.post(f"/inventare/{inventar.id}/regenerare")
    assert inventar.stare == "GENERAT"
    client.post(f"/inventare/{inventar.id}/stergere")
    assert session.scalar(select(Inventar)) is None


def test_consultarea_doar_vede(client, lume, creeaza):
    client.post("/inventare/nou", data={"data": "2025-11-28", "confirmat": "1"})
    client.post("/deconectare")
    creeaza("cons", rol="CONSULTARE")
    client.post("/autentificare", data={"username": "cons", "parola": "parola123"})
    assert client.get("/inventare").status_code == 200
    assert client.get("/inventare/1").status_code == 200
    assert client.get("/inventare/nou").status_code == 403
    assert client.post("/inventare/1/regenerare").status_code == 403


def test_tiparire_din_pagina_inventarului(client, session, lume):
    from io import BytesIO

    from pypdf import PdfReader

    client.post("/inventare/nou", data={"data": "2025-11-28", "confirmat": "1"})
    inventar = session.scalar(select(Inventar))
    pagina = client.get(f"/inventare/{inventar.id}").get_data(as_text=True)
    assert "Declarații de sfârșit OI" in pagina and "Liste MF" in pagina

    r = client.get(f"/inventare/{inventar.id}/pdf/liste/MF")
    assert r.status_code == 200 and r.mimetype == "application/pdf"
    assert "Inventar_2025-11-28_tiraj_MF.pdf" in r.headers["Content-Disposition"]
    assert "200,00" in PdfReader(BytesIO(r.data)).pages[0].extract_text()

    r = client.get(f"/inventare/{inventar.id}/pdf/inceput/OI?persoana_id={lume['a'].id}")
    assert "HRITCU_MARCELA_OI" in r.headers["Content-Disposition"]

    inventar.stare = "NEACTUALIZAT"
    session.commit()
    r = client.get(f"/inventare/{inventar.id}/pdf/liste/MF", follow_redirects=True)
    assert "regenerați-l înainte de tipărire" in r.get_data(as_text=True)


def test_sabloane_doar_admin(client, session, lume):
    from inventar.domain.sabloane import citeste_sablon

    assert client.get("/admin/sabloane").status_code == 403
    client.post("/deconectare")
    client.post("/autentificare", data={"username": "admin", "parola": "parola123"})
    assert "(11) Dacă mai aveţi" in client.get("/admin/sabloane").get_data(as_text=True)
    client.post("/admin/sabloane", data={"cod": "declaratie_inceput", "corp": "(1) Text nou\r\n(2) Altul",
                                         "actiune": "salveaza"})
    assert citeste_sablon(session, "declaratie_inceput")[1] == "(1) Text nou\n(2) Altul"
    client.post("/admin/sabloane", data={"cod": "declaratie_inceput", "actiune": "reseteaza"})
    assert "(11)" in citeste_sablon(session, "declaratie_inceput")[1]


def test_pagina_listari(client, session, lume):
    from io import BytesIO

    from openpyxl import load_workbook

    client.post("/inventare/nou", data={"data": "2025-11-28", "confirmat": "1"})
    inventar = session.scalar(select(Inventar))
    url = f"/inventare/{inventar.id}/listari"
    pagina = client.get(url).get_data(as_text=True)
    assert "Filtru curent: Tot" in pagina and "HRITCU MARCELA" in pagina

    # Filtrul se transmite în toate legăturile (parametri repetați pentru selecții multiple).
    pagina = client.get(f"{url}?categorii=OI&conturi=303&centralizator=categorii").get_data(as_text=True)
    assert "Categorii: OI; Conturi: 303" in pagina
    assert "export.xlsx?categorii=OI&amp;conturi=303" in pagina
    assert "TOTAL GENERAL" in pagina and "10,00" in pagina

    r = client.get(f"/inventare/{inventar.id}/export.xlsx?categorii=MF")
    assert r.status_code == 200 and "attachment" in r.headers["Content-Disposition"]
    assert len(list(load_workbook(BytesIO(r.data)).active.iter_rows())) == 3  # antet + 2 MF

    r = client.get(f"/inventare/{inventar.id}/centralizator/locatii.pdf")
    assert r.mimetype == "application/pdf"
    r = client.get(f"/inventare/{inventar.id}/centralizator/necunoscut.pdf", follow_redirects=True)
    assert "necunoscut" in r.get_data(as_text=True)

    r = client.get(f"/inventare/{inventar.id}/pdf/liste/MF?categorii=OI", follow_redirects=True)
    assert "Nu există bunuri" in r.get_data(as_text=True)
