"""Export / import al adreselor de e-mail ale persoanelor, prin Excel.

Se exportă gestiunile cu e-mailul curent; fișierul completat se încarcă, se arată ce se
schimbă și abia la confirmare se salvează. Potrivirea se face după ID (nu după nume), iar
importul atinge numai câmpul e-mail: numele, titularul și starea rămân cum sunt.
"""
from dataclasses import dataclass, field
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from inventar import jurnal
from inventar.domain.comun import EroareValidare
from inventar.domain.persoane import adresa_valida
from inventar.drepturi import Actor, cere_rol
from inventar.models import Persoana

COLOANE = ("ID", "Gestiune", "Titular", "E-mail", "Activ")
_ANTET_ID = {"id"}
_ANTET_EMAIL = {"e-mail", "email", "e-mail curent"}


@dataclass(frozen=True)
class Schimbare:
    persoana_id: int
    nume: str
    vechi: str | None
    nou: str
    rand: int


@dataclass
class Analiza:
    schimbari: list[Schimbare] = field(default_factory=list)
    neschimbate: int = 0
    goale: int = 0  # celulă goală = lăsat neschimbat, nu ștergere
    erori: list[str] = field(default_factory=list)
    avertismente: list[str] = field(default_factory=list)


def exporta(session: Session, actor: Actor, include_inactive: bool = False) -> bytes:
    cere_rol(actor, "CONSULTARE")
    interogare = select(Persoana).order_by(Persoana.nume)
    if not include_inactive:
        interogare = interogare.where(Persoana.activ.is_(True))
    persoane = session.scalars(interogare).all()

    carte = Workbook()
    foaie = carte.active
    foaie.title = "E-mailuri"
    foaie.append(list(COLOANE))
    for celula in foaie[1]:
        celula.font = Font(bold=True)
    gri = PatternFill("solid", fgColor="E7E6E6")
    for p in persoane:
        foaie.append([p.id, p.nume, p.titular.nume if p.titular else "", p.email or "", "da" if p.activ else "nu"])
        # ID-ul e cheia de potrivire la import; coloanele gri nu se modifică.
        for coloana in (1, 2, 3, 5):
            foaie.cell(foaie.max_row, coloana).fill = gri
    foaie.freeze_panes = "A2"
    foaie.auto_filter.ref = foaie.dimensions
    for litera, latime in zip("ABCDE", (8, 42, 36, 40, 8)):
        foaie.column_dimensions[litera].width = latime

    jurnal.scrie(session, actor, "EXPORT_EMAIL_PERSOANE", detalii={"randuri": len(persoane),
                                                                    "inclusiv_inactive": include_inactive})
    session.commit()
    continut = BytesIO()
    carte.save(continut)
    return continut.getvalue()


def _intreg(valoare: object) -> int | None:
    try:
        numar = float(str(valoare).strip().replace(",", "."))
    except ValueError:
        return None
    return int(numar) if numar == int(numar) and numar > 0 else None


def _randuri(continut: bytes) -> list[tuple[int, object, object]]:
    """(număr rând, ID, e-mail) pentru fiecare rând nevid; coloanele se caută după antet."""
    try:
        carte = load_workbook(BytesIO(continut), read_only=True, data_only=True)
        randuri = list(carte.worksheets[0].iter_rows(values_only=True))
    except Exception as e:  # openpyxl ridică tipuri felurite pe fișiere deteriorate
        raise EroareValidare("Fișierul nu este un document Excel (.xlsx) valid.") from e
    if not randuri:
        raise EroareValidare("Fișierul este gol.")
    antet = [str(c).strip().lower() if c is not None else "" for c in randuri[0]]
    col_id = next((i for i, c in enumerate(antet) if c in _ANTET_ID), None)
    col_email = next((i for i, c in enumerate(antet) if c in _ANTET_EMAIL), None)
    if col_id is None or col_email is None:
        raise EroareValidare("Fișierul trebuie să aibă coloanele „ID” și „E-mail” pe primul rând "
                             "(folosiți fișierul exportat din aplicație).")
    rezultat = []
    for nr, rand in enumerate(randuri[1:], start=2):
        rand = tuple(rand) + (None,) * (max(col_id, col_email) + 1 - len(rand))
        if all(c is None or str(c).strip() == "" for c in rand):
            continue
        rezultat.append((nr, rand[col_id], rand[col_email]))
    return rezultat


def analizeaza(session: Session, actor: Actor, continut: bytes) -> Analiza:
    """Compară fișierul cu baza; nu modifică nimic."""
    cere_rol(actor, "OPERATOR")
    persoane = {p.id: p for p in session.scalars(select(Persoana))}
    analiza = Analiza()
    vazute: set[int] = set()
    for nr, id_brut, email_brut in _randuri(continut):
        persoana_id = _intreg(id_brut)
        if persoana_id is None:
            analiza.erori.append(f"Rândul {nr}: ID lipsă sau necitibil.")
            continue
        persoana = persoane.get(persoana_id)
        if persoana is None:
            analiza.erori.append(f"Rândul {nr}: nu există nicio persoană cu ID {persoana_id}.")
            continue
        if persoana_id in vazute:
            analiza.erori.append(f"Rândul {nr}: ID {persoana_id} („{persoana.nume}”) apare de două ori în fișier.")
            continue
        vazute.add(persoana_id)
        email = str(email_brut).strip() if email_brut is not None else ""
        if not email:
            analiza.goale += 1
        elif not adresa_valida(email):
            analiza.erori.append(f"Rândul {nr}: adresa „{email}” a persoanei „{persoana.nume}” nu este validă.")
        elif (persoana.email or "").lower() == email.lower():
            analiza.neschimbate += 1
        else:
            analiza.schimbari.append(Schimbare(persoana_id, persoana.nume, persoana.email, email, nr))
    _avertizeaza_dubluri(analiza, persoane)
    return analiza


def _avertizeaza_dubluri(analiza: Analiza, persoane: dict[int, Persoana]) -> None:
    """Aceeași adresă la persoane fizice diferite e probabil o greșeală (gestiunile aceluiași titular nu contează)."""
    dupa_adresa: dict[str, dict[int, str]] = {}
    noi = {s.persoana_id: s.nou for s in analiza.schimbari}
    for p in persoane.values():
        email = noi.get(p.id, p.email)
        if email:
            dupa_adresa.setdefault(email.lower(), {})[p.titular_id or p.id] = p.nume
    for s in analiza.schimbari:
        grupuri = dupa_adresa[s.nou.lower()]
        if len(grupuri) > 1:
            analiza.avertismente.append(f"Adresa {s.nou} este folosită de mai multe persoane: "
                                        f"{', '.join(sorted(grupuri.values()))}.")


def aplica(session: Session, actor: Actor, schimbari: dict[int, str]) -> int:
    """Atomic: dacă o adresă sau o persoană nu e în regulă, nu se salvează nimic."""
    cere_rol(actor, "OPERATOR")
    efectuate = []
    for persoana_id, email in schimbari.items():
        persoana = session.get(Persoana, persoana_id)
        email = email.strip()
        if persoana is None:
            raise EroareValidare(f"Persoana cu ID {persoana_id} nu există.")
        if not adresa_valida(email):
            raise EroareValidare(f"Adresa „{email}” nu este validă.")
        if persoana.email != email:
            efectuate.append({"persoana_id": persoana.id, "nume": persoana.nume, "vechi": persoana.email,
                              "nou": email})
            persoana.email = email
    # Adresa se citește la trimitere, nu e în instantaneu: niciun inventar nu devine NEACTUALIZAT.
    jurnal.scrie(session, actor, "EMAIL_PERSOANE_IMPORTAT", detalii={"nr": len(efectuate), "schimbari": efectuate})
    session.commit()
    return len(efectuate)
