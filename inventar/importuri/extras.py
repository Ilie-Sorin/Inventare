"""Citirea extrasului din aplicația existentă (Imp_Ap_Existenta.xlsx), folosit o singură dată (§9.2)."""
from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook

from inventar.domain.comun import curata_text
from inventar.domain.normalizare import nr_inventar, nr_inventar_optional
from inventar.importuri.mf_baza import RezultatCitire

COLOANE = ("Tip", "Locatia", "Activitatea", "Persoana", "NrInventar", "NrInventarVechi")
TIP_CATEGORIE = {"MijloaceFixe": "MF", "ObiecteInventar": "OI"}


@dataclass
class RandExtras:
    rand: int
    categorie: str
    locatie: str
    activitate: str
    persoana: str
    nr_inventar: str
    nr_inventar_vechi: str | None


def citeste_extras(continut: bytes, fisier: str) -> RezultatCitire:
    try:
        carte = load_workbook(BytesIO(continut), read_only=True, data_only=True)
    except Exception:
        return RezultatCitire([], [f"{fisier}: nu este un fișier Excel .xlsx valid."])
    foaie = carte.worksheets[0]
    randuri = foaie.iter_rows(values_only=True)
    antet = [curata_text(str(v)) if v is not None else "" for v in next(randuri, ())]
    lipsa = [c for c in COLOANE if c not in antet]
    if lipsa:
        return RezultatCitire([], [f"{fisier}: lipsesc coloanele {', '.join(lipsa)}."])
    index = {c: antet.index(c) for c in COLOANE}

    rezultat = RezultatCitire([], [])
    for numar, valori in enumerate(randuri, start=2):
        v = {c: (valori[i] if i < len(valori) else None) for c, i in index.items()}
        if all(x in (None, "") for x in v.values()):
            continue
        try:
            tip = curata_text(str(v["Tip"] or ""))
            if tip not in TIP_CATEGORIE:
                raise ValueError(f"Tip „{tip}” necunoscut (MijloaceFixe sau ObiecteInventar)")
            campuri = {c: curata_text(str(v[c] or "")) for c in ("Locatia", "Activitatea", "Persoana")}
            for camp, valoare in campuri.items():
                if not valoare:
                    raise ValueError(f"{camp} lipsă")
            rezultat.randuri.append(RandExtras(
                rand=numar,
                categorie=TIP_CATEGORIE[tip],
                locatie=campuri["Locatia"],
                activitate=campuri["Activitatea"],
                persoana=campuri["Persoana"],
                nr_inventar=nr_inventar(v["NrInventar"]),
                nr_inventar_vechi=nr_inventar_optional(v["NrInventarVechi"]),
            ))
        except ValueError as e:
            rezultat.erori.append(f"{fisier}, rândul {numar}: {e}.")
    carte.close()
    return rezultat
