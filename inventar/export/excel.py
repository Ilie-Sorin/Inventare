"""Export Excel (.xlsx) și CSV: linii de inventar și centralizatoare."""
import csv
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from inventar.domain.centralizatoare import RandCentralizator
from inventar.formatare import data as data_ro
from inventar.formatare import suma

COLOANE_LINII = (
    ("nr_inventar", "Nr. inventar"), ("nr_inventar_vechi", "Nr. inventar vechi"), ("denumire", "Denumire"),
    ("serie", "Serie"), ("um", "U/M"), ("categorie", "Categorie"), ("cont", "Cont"),
    ("cont_sintetic", "Cont sintetic"), ("pret", "Valoare"), ("data_amortizare_integrala", "Data amortizării integrale"),
    ("gestiune", "Gestiune"), ("titular", "Titular"), ("locatie", "Locație"), ("activitate", "Activitate"),
)
FORMAT_BANI = "#,##0.00"
FORMAT_DATA = "DD.MM.YYYY"


def _salveaza(carte: Workbook) -> bytes:
    continut = BytesIO()
    carte.save(continut)
    return continut.getvalue()


def _latimi(foaie, latimi: Iterable[int]) -> None:
    for i, latime in enumerate(latimi, start=1):
        foaie.column_dimensions[get_column_letter(i)].width = latime


def linii_xlsx(linii: Iterable, titlu: str) -> bytes:
    """Un rând pe bun; valorile rămân numere și date reale, ca să se poată calcula în Excel."""
    carte = Workbook()
    foaie = carte.active
    foaie.title = "Bunuri"
    foaie.append([eticheta for _, eticheta in COLOANE_LINII])
    for celula in foaie[1]:
        celula.font = Font(bold=True)
    for linie in linii:
        foaie.append([getattr(linie, camp) for camp, _ in COLOANE_LINII])
        rand = foaie.max_row
        foaie.cell(rand, 9).number_format = FORMAT_BANI
        foaie.cell(rand, 10).number_format = FORMAT_DATA
    foaie.freeze_panes = "A2"
    foaie.auto_filter.ref = foaie.dimensions
    _latimi(foaie, (12, 12, 45, 14, 6, 9, 14, 9, 13, 13, 30, 30, 22, 22))
    carte.properties.title = titlu
    return _salveaza(carte)


def _text_csv(valoare: object) -> str:
    # Pentru Excel în limba română: virgulă zecimală, dată dd.mm.yyyy.
    if isinstance(valoare, Decimal):
        return suma(valoare, separator_mii=False)
    if isinstance(valoare, date):
        return data_ro(valoare)
    return "" if valoare is None else str(valoare)


def linii_csv(linii: Iterable) -> bytes:
    """Separator „;” și BOM UTF-8, ca Excel-ul românesc să deschidă corect diacriticele."""
    text = StringIO()
    scriitor = csv.writer(text, delimiter=";", lineterminator="\r\n")
    scriitor.writerow([eticheta for _, eticheta in COLOANE_LINII])
    for linie in linii:
        scriitor.writerow([_text_csv(getattr(linie, camp)) for camp, _ in COLOANE_LINII])
    return text.getvalue().encode("utf-8-sig")


def centralizator_xlsx(randuri: list[RandCentralizator], titlu: str, subtitlu: str) -> bytes:
    carte = Workbook()
    foaie = carte.active
    foaie.title = "Centralizator"
    foaie.append([titlu])
    foaie["A1"].font = Font(bold=True, size=13)
    foaie.append([subtitlu])
    foaie.append([])
    foaie.append(["", "Număr bunuri", "Valoare"])
    for celula in foaie[4]:
        celula.font = Font(bold=True)
    for r in randuri:
        foaie.append(["   " * r.nivel + r.eticheta, r.numar, r.valoare])
        rand = foaie.max_row
        foaie.cell(rand, 3).number_format = FORMAT_BANI
        if r.tip != "detaliu":
            for coloana in (1, 2, 3):
                foaie.cell(rand, coloana).font = Font(bold=True)
    _latimi(foaie, (55, 14, 18))
    return _salveaza(carte)


def tabele_xlsx(foi: list[tuple[str, list[str], list[list]]], titlu: str) -> bytes:
    """Câte o foaie pe tabel: (nume foaie, antet, rânduri). Sumele (Decimal) primesc format de bani."""
    carte = Workbook()
    carte.remove(carte.active)
    for nume, antet, randuri in foi:
        foaie = carte.create_sheet(nume[:31])
        foaie.append(antet)
        for celula in foaie[1]:
            celula.font = Font(bold=True)
        for rand in randuri:
            foaie.append(rand)
            for i, valoare in enumerate(rand, start=1):
                if isinstance(valoare, Decimal):
                    foaie.cell(foaie.max_row, i).number_format = FORMAT_BANI
        foaie.freeze_panes = "A2"
        _latimi(foaie, [max(12, min(80, max((len(str(r[i])) for r in [antet, *randuri]), default=10) + 2))
                        for i in range(len(antet))])
    carte.properties.title = titlu
    return _salveaza(carte)
